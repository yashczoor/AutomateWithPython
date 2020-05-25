#! python3
#will create an excel file, write data to it and insert bar chart

import openpyxl
from copy import deepcopy #to copy charts


wb = openpyxl.Workbook()
ws = wb.active

rows = [
    ('Wojewodztwo','Wartosc'),
    ('Mazowieckie',100),
    ('Wielkopolskie',200),
    ('Łódzkie',300),
    ('Lubelskie',400)
    ]

for row in rows:
    ws.append(row)

chart = openpyxl.chart.BarChart() #parenthesis are important
chart.type = 'col'
chart.title = 'Bar Chart'
chart.y_axis.title = 'Value'
chart.x_axis.title = 'Item'

data = openpyxl.chart.Reference(ws,2,1,2,5) #VALUES
items = openpyxl.chart.Reference(ws,min_col=1,min_row=2,max_row=5) #LABELS

chart.add_data(data,titles_from_data=True)
chart.set_categories(items)
chart.shape = 4
ws.add_chart(chart, "A7")

chart2 = deepcopy(chart)
chart2.style = 11
chart2.type = "bar"
chart2.title = "Horizontal Bar Chart"

ws.add_chart(chart2, "K7")

wb.save('D://barChart.xlsx')
